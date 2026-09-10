#!/usr/bin/env python3
"""Build the materialized-input presentation from saved evidence; never run queries.

Run with the presentation environment (python-pptx and matplotlib required).
Original deck, notes, experiment artifacts, and historical ORQ figures are preserved.
"""
from pathlib import Path
import csv
import hashlib
import json
import re
from collections import Counter

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

ROOT = Path(__file__).resolve().parents[1]
RUN = ROOT / 'results/duckdb-tpch-selectivity/20260908T200506.888858Z-sf0p1'
OUT = ROOT / 'outputs/duckdb-materialized-presentation'
DECK = ROOT / 'duckdb-orq-materialized-inputs-10min.pptx'
NOTES = ROOT / 'duckdb-orq-materialized-inputs-speaker-notes.md'
ORIGINAL = ROOT / 'duckdb-orq-verified-metrics-10min.pptx'
ORIGINAL_NOTES = ROOT / 'duckdb-orq-verified-metrics-speaker-notes.md'
COLORS = {1:'#0072B2',3:'#D55E00',5:'#009E73',8:'#AA4499',9:'#B8860B'}
INK = '171717'
GRAY = '5F6368'
ORANGE = 'D55E00'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def evidence():
    manifest = json.loads((RUN/'manifest.json').read_text())
    assert manifest['status'] == 'complete'
    rows = list(csv.DictReader((RUN/'points.csv').open()))
    materialized = [r for r in rows if r['input_mode']=='materialized']
    assert all(r['result_equivalent']=='True' for r in materialized)
    selected = [r for r in materialized if r['control']=='optimized']
    lookup = {(int(r['query_id']),int(r['selected_rows'])):r for r in selected}
    common = [r for r in selected if r['point_kind']=='common']
    assert Counter(r['query_id'] for r in common)=={str(q):31 for q in COLORS}
    samples = {}
    for r in csv.DictReader((RUN/'samples.csv').open()):
        samples.setdefault(r['point_id'],[]).append(float(r['seconds']))
    import statistics
    for r in common:
        vals = samples[r['point_id']]
        assert len(vals)==30
        assert abs(statistics.median(vals)-float(r['median_seconds']))<1e-12
    transitions=[]
    for search in json.loads((RUN/'transition_search.json').read_text()):
        if search['input_mode']!='materialized': continue
        q=search['query_id']
        for i,t in enumerate(search['transitions'],1):
            a,b=(lookup[q,t[f'{side}_count']] for side in ('low','high'))
            changed=[d for d in ('association','orientation','physical') if a[d+'_plan_id']!=b[d+'_plan_id']]
            assert changed==t['changed_dimensions']
            assert t['adjacent_rows']==(t['high_count']-t['low_count']==1)
            for r in (a,b): assert (RUN/r['raw_dir']/'explain.json').is_file()
            entry={'query':q,'transition':f'Q{q}-T{i}',**t,
                   'low_percent':float(a['actual_percent']),'high_percent':float(b['actual_percent'])}
            for side,r in [('low',a),('high',b)]:
                entry[side+'_point_id']=r['point_id']
                entry[side+'_plan_path']=r['raw_dir']+'/explain.json'
                for d in ('association','orientation','physical'):entry[side+'_'+d+'_id']=r[d+'_plan_id']
            transitions.append(entry)
    return manifest,rows,selected,common,lookup,transitions


def chart(common,transitions):
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,
                         'axes.spines.right':False,'axes.labelcolor':'#333333'})
    fig=plt.figure(figsize=(12.4,5.7),facecolor='white')
    gs=fig.add_gridspec(2,2,width_ratios=[1.75,1],height_ratios=[2.4,1.6],
                        left=.066,right=.985,bottom=.14,top=.89,wspace=.13,hspace=.18)
    axes=[fig.add_subplot(gs[0,i]) for i in range(2)]
    tracks=[fig.add_subplot(gs[1,i],sharex=axes[i]) for i in range(2)]
    for q,c in COLORS.items():
        data=sorted((r for r in common if int(r['query_id'])==q),key=lambda r:float(r['actual_percent']))
        x=[float(r['actual_percent']) for r in data]
        y=[1000*float(r['median_seconds']) for r in data]
        for ax in axes:
            ax.plot(x,y,'o-',color=c,lw=1.65,ms=2.8,label=f'Q{q}',zorder=3)
            ax.fill_between(x,[1000*float(r['median_ci_low']) for r in data],
                            [1000*float(r['median_ci_high']) for r in data],color=c,alpha=.13,lw=0)
    axes[0].set_xlim(0,100);axes[1].set_xlim(0,10)
    axes[0].set_title('Full input range',loc='left',fontsize=12,fontweight='bold')
    axes[1].set_title('Zoom: 0–10% retained input',loc='left',fontsize=12,fontweight='bold')
    for ax in axes:
        ax.set_ylim(bottom=0);ax.grid(alpha=.18);ax.tick_params(labelbottom=False)
    # Same runtime scale keeps the inset directly comparable.
    axes[1].set_ylim(axes[0].get_ylim())
    axes[0].set_ylabel('Median query time (ms)')
    fig.legend(*axes[0].get_legend_handles_labels(),loc='upper center',bbox_to_anchor=(.49,1.015),
               ncol=5,frameon=False,columnspacing=2.8)
    positions={q:4-i for i,q in enumerate(COLORS)}
    for ax in tracks:
        for q,y in positions.items(): ax.axhline(y,color=COLORS[q],alpha=.24,lw=.8)
        ax.set_yticks(list(positions.values()),[f'Q{q}' for q in positions])
        for label,q in zip(ax.get_yticklabels(),COLORS): label.set_color(COLORS[q])
        ax.set_ylim(-.6,4.6);ax.tick_params(axis='y',length=0)
        ax.set_xlabel('Actual retained input (%)',fontsize=10)
        ax.spines['left'].set_visible(False)
        ax.grid(axis='x',alpha=.12)
        ax.text(.012,positions[1]+.1,'no observed changes',color=COLORS[1],fontsize=8,
                transform=ax.get_yaxis_transform())
    # Structural changes also change the physical fingerprint. Reserve the square
    # for physical-only changes to avoid presenting this as an independent event.
    for t in transitions:
        q=t['query'];c=COLORS[q];x=t['high_percent'];y=positions[q];dims=t['changed_dimensions']
        for col,ax in enumerate(tracks):
            if col==1 and t['low_percent']>10:continue
            if not t['adjacent_rows']: ax.axvspan(t['low_percent'],x,color=c,alpha=.12)
            if 'association' in dims: ax.scatter([x],[y+.16],marker='D',color=c,s=24,zorder=4)
            if 'orientation' in dims: ax.scatter([x],[y-.12],marker='v',color=c,s=22,zorder=4)
            if dims==['physical']: ax.scatter([x],[y],marker='s',color=c,s=23,zorder=4)
            # Label association thresholds in the panel where they are legible.
            if 'association' in dims and ((col==0 and x>10) or (col==1 and x<=10)):
                dx,dy=0,10
                if q==8 and 4<x<4.2: dx,dy=-17,12
                if q==8 and 4.2<x<4.4: dx,dy=17,12
                ax.annotate(f'{x:.3f}%',(x,y+.16),xytext=(dx,dy),textcoords='offset points',
                            fontsize=7,color=c,ha='center',arrowprops={'arrowstyle':'-','color':c,'lw':.5})
    handles=[Line2D([],[],color='#555',marker=m,linestyle='',markersize=5,label=l) for m,l in
             [('D','Join association'),('v','Build/probe orientation'),('s','Physical-only change')]]
    fig.legend(handles=handles,loc='lower center',bbox_to_anchor=(.53,-.005),ncol=3,frameon=False,fontsize=9)
    for ext in ['png','svg','pdf']:
        fig.savefig(OUT/f'duckdb-materialized-lines.{ext}',dpi=220)
    plt.close(fig)


def text(slide,value,x,y,w,h,size=20,bold=False,color=INK,align=PP_ALIGN.LEFT):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=shape.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
        p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        p.alignment=align;p.space_after=Pt(5)
    return shape


def replace_shape(slide,name,value):
    shape=next(s for s in slide.shapes if s.name==name)
    # Retain XML paragraph/run styling rather than assigning TextFrame.text.
    ts=shape._element.xpath('.//a:t')
    ts[0].text=value
    for t in ts[1:]: t.text=''


def blank_body(slide,title,source):
    for shape in list(slide.shapes):
        if shape.name.endswith('-title') or shape.name.endswith('-source') or shape.name.endswith('-number'):
            if shape.name.endswith('-title'):replace_shape(slide,shape.name,title)
            if shape.name.endswith('-source'):replace_shape(slide,shape.name,source)
        else:shape._element.getparent().remove(shape._element)


def band(slide,value,y=6.35):
    s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.65), Inches(y), Inches(12.0), Inches(.43))
    s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string('F2F2F2');s.line.fill.background()
    text(slide,value,.8,y+.075,11.7,.30,15,True,align=PP_ALIGN.CENTER)


def table(slide,headers,rows,widths,y=2.0,row_h=.66,size=19):
    x=.7
    for label,w in zip(headers,widths):text(slide,label,x,y,w-.12,.45,16,True,GRAY);x+=w
    for i,row in enumerate(rows):
        yy=y+.55+i*row_h
        if i%2==0:
            s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.65), Inches(yy-.06), Inches(12), Inches(row_h-.02))
            s.fill.solid();s.fill.fore_color.rgb=RGBColor.from_string('F5F5F5');s.line.fill.background()
        x=.7
        for j,(label,w) in enumerate(zip(row,widths)):
            text(slide,label,x,yy,w-.16,row_h-.07,size,j==0);x+=w


def parse_tree(s):
    tokens=re.findall(r'join|[a-z_0-9]+|[(),]',s);at=0
    def walk():
        nonlocal at
        token=tokens[at];at+=1
        if token!='join':return token
        assert tokens[at]=='(';at+=1
        a=walk();assert tokens[at]==',';at+=1
        b=walk();assert tokens[at]==')';at+=1
        return (a,b)
    root=walk();assert at==len(tokens);return root


def tree(slide,association,x,y,w,h):
    root=parse_tree(association)
    def leaves(n):return 1 if isinstance(n,str) else leaves(n[0])+leaves(n[1])
    def depth(n):return 0 if isinstance(n,str) else 1+max(depth(n[0]),depth(n[1]))
    count=leaves(root);maxdepth=depth(root);leaf_index=0
    nodes=[];edges=[]
    labels={'nation_region':'Nation\n(region)','nation_name':'Nation\n(name)'}
    def layout(n,level):
        nonlocal leaf_index
        if isinstance(n,str):
            xx=x+w*(leaf_index+.5)/count;leaf_index+=1;yy=y+h
            nodes.append((xx,yy,labels.get(n,n.title()),False));return xx,yy
        a=layout(n[0],level+1);b=layout(n[1],level+1)
        pos=((a[0]+b[0])/2,y+level*(h-.45)/maxdepth)
        edges.extend([(pos,a),(pos,b)]);nodes.append((*pos,'⋈',True));return pos
    layout(root,0)
    for (a,b),(c,d) in edges:
        line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(a),Inches(b+.13),Inches(c),Inches(d))
        line.line.color.rgb=RGBColor.from_string('8B9298');line.line.width=Pt(1.2)
    for xx,yy,label,is_join in nodes:
        text(slide,label,xx-(.21 if is_join else .43),yy-.08,.42 if is_join else .86,.45 if is_join else .65,
             22 if is_join else 11,True,ORANGE if is_join else INK,PP_ALIGN.CENTER)


def build_slides(lookup,transitions,selected):
    p=Presentation(ORIGINAL)
    replace_shape(p.slides[1],'mini-insight','How do smaller physical inputs change DuckDB’s selected join structure?')
    s=p.slides[2];blank_body(s,'Q5 varies a materialized Orders input','8 Sep 2026 · saved Q5 materialized plans · association only; build/probe orientation omitted')
    text(s,'Rank once  →  retain k Orders  →  materialize + ANALYZE  →  execute Q5',.7,1.3,12,.45,22,True)
    text(s,'Replace the date window with the selected input; keep the ASIA region predicate.',.7,1.94,12,.4,18,color=GRAY)
    for x,k,label in [(.7,1500,'1% retained · 1,500 Orders'),(7.0,3000,'2% retained · 3,000 Orders')]:
        text(s,label,x,2.57,5.7,.4,21,True)
        tree(s,lookup[5,k]['association_fingerprint'],x,3.17,5.5,2.0)
    band(s,'Observed association boundary: 2,501 → 2,502 Orders (1.668% at the new plan)')
    s=p.slides[3];blank_body(s,'DuckDB data and controlled input sizes','DuckDB 1.5.4 · TPC-H DBGEN SF 0.1 · one fixed database · seeded nested subsets')
    text(s,'Retention = selected rows / original rows in the varied relation',.7,1.35,12,.45,23,True)
    table(s,['Query','Materialized input','Original rows','Predicate replaced'],[
        ['Q1','Lineitem','600,572','Ship-date cutoff'],['Q3','Lineitem','600,572','Ship-date condition only'],
        ['Q5','Orders','150,000','Order-date window'],['Q8','Part','20,000','Part-type filter'],
        ['Q9','Part','20,000','Part-name LIKE filter']], [1.0,2.5,2.3,6.2],size=20)
    text(s,'31 common targets: 0.01–100%  ·  additional probes locate observed plan changes',.7,6.0,12,.3,17,color=GRAY)
    band(s,'Stable seeded rank ≤ k; subsets are nested. Other query predicates remain fixed.')
    s=p.slides[4];blank_body(s,'DuckDB execution and measurement','8 Sep 2026 · manifest.json, points.csv, samples.csv · completed materialized-input study')
    text(s,'1 thread  ·  5 warmups  ·  30 timed repetitions  ·  separate operator profile',.7,1.35,12,.45,22,True)
    table(s,['Stage','Recorded method'],[
        ['Prepare','Materialize selected rows; ANALYZE the selected table'],
        ['Compare','Optimized + two fixed orders for joins; Q1 optimized only'],
        ['Time','Execute through fetchall; randomized point/control ordering'],
        ['Summarize','Median + 95% bootstrap median interval (2,000 resamples)'],
        ['Validate','2,026 materialized point/control records pass result checks']], [2.0,10.0],y=2.0,row_h=.64,size=19)
    text(s,'Fixed orders disable join_order and build_side_probe_side; other optimizations remain.',.7,5.85,12,.35,16,color=GRAY)
    band(s,'Query timings exclude materialization, ANALYZE and profiling.')
    s=p.slides[5];blank_body(s,'DuckDB runtime and observed plan changes','Materialized inputs · Q1/Q3 Lineitem; Q5 Orders; Q8/Q9 Part · 8 Sep 2026 · query-only timings')
    s.shapes.add_picture(str(OUT/'duckdb-materialized-lines.png'),Inches(.44),Inches(1.13),width=Inches(12.45),height=Inches(5.72))
    text(s,'Lines: 31 measured medians/query + 95% intervals. Tracks: refinement evidence; search not exhaustive.',.7,6.83,12,.23,11,color=GRAY)
    s=p.slides[6];blank_body(s,'Q8 changes its selected join association','8 Sep 2026 · exact saved associations at 1% and 5% · join subtrees; physical orientation omitted')
    for x,k,label in [(.65,200,'1% retained · 200 Parts'),(6.98,1000,'5% retained · 1,000 Parts')]:
        text(s,label,x,1.4,5.8,.4,23,True)
        tree(s,lookup[8,k]['association_fingerprint'],x,2.15,5.65,3.05)
    text(s,'At 1%: Part joins Lineitem early.',.7,5.97,5.7,.3,18)
    text(s,'At 5%: the region/customer branch joins first.',6.98,5.97,5.7,.3,17)
    band(s,'This motivates ORQ candidates; it does not establish an ORQ speedup.')
    replace_shape(p.slides[7],'slide-4-title','Earlier ORQ study: data and measurement')
    replace_shape(p.slides[10],'mini-insight','ORQ uses the canonical transfer frozen on 1 Sep, before the materialized study.')
    replace_shape(p.slides[11],'slide-4-title','Earlier ORQ runtime: original and transferred plans')
    replace_shape(p.slides[13],'slide-5-source','Earlier fixed-predicate ORQ evidence · new DuckDB input-size study motivates further tests')
    s=p.slides[16];blank_body(s,'Backup · Observed materialized plan transitions','Source: transition_search.json + endpoint plans · A = association; O = orientation; P = physical-only')
    text(s,'New-plan thresholds (% of original input); all located between adjacent row counts',.7,1.33,12,.4,21,True)
    rows=[]
    for q in COLORS:
        ts=[t for t in transitions if t['query']==q]
        labels=[]
        for t in ts:
            dims=t['changed_dimensions']; code=('A' if 'association' in dims else '')+('O' if 'orientation' in dims else '')
            code=code or 'P';labels.append(f"{t['high_percent']:.3f} {code}")
        value='None observed (no joins)' if not labels else '   ·   '.join(labels[:4])+ ('\n'+'   ·   '.join(labels[4:]) if len(labels)>4 else '')
        count=len({r['association_plan_id'] for r in selected if int(r['query_id'])==q})
        rows.append([f'Q{q}',str(count),value])
    table(s,['Query','Associations','New-plan retained-input % and change type'],rows,[1,1.8,9.2],y=2.0,row_h=.73,size=16)
    text(s,'Direct-input arm: one observed association per query. Materialization changes the planning input.',.7,6.29,12,.3,16,color=GRAY)
    text(s,'Rounded percentages here; exact counts and all plan IDs are in the notes and transition ledger.',.7,6.7,12,.25,13,color=GRAY)
    return p


SCRIPTS = {
2:('A selective filter can leave secure work unchanged','0:35–1:10',
'Suppose a filter accepts one hundred of a thousand rows. Plaintext operators can process those hundred survivors. In ORQ, validity bits remain secret and the table can still occupy a thousand positions. Later operators may process that physical length. These numbers are illustrative. To study how physical input size affects planning, I now materialize controlled subsets in DuckDB. This exposes smaller inputs to its optimizer; it does not demonstrate that ORQ can reveal or compact private survivors.'),
3:('Q5 varies a materialized Orders input','1:10–2:05',
'Q5 makes the new experiment concrete. I rank Orders once with a seeded deterministic ordering, retain the first k rows, materialize that subset, and analyze its statistics before executing the query. This replaces the original order-date window; the ASIA region condition remains. One percent retains fifteen hundred orders; two percent retains three thousand. Read both trees from the leaves upward. The selected association differs: at one percent Customer joins Orders before the region branch, while at two percent Customer joins the region branch first. Refinement locates this association change between two thousand five hundred one and two thousand five hundred two orders.'),
4:('DuckDB data and controlled input sizes','2:05–2:45',
'The study uses one fixed TPC-H DBGEN database at scale factor zero point one. I vary Lineitem for Q1 and Q3, Orders for Q5, and Part for Q8 and Q9. Retention means selected rows divided by the original rows in that relation. All queries share thirty-one percentage targets, from zero point zero one to one hundred percent, with row-count rounding. Seeded ranks make the subsets nested and reproducible. The listed predicate is replaced; other query conditions stay fixed. These are controlled query structures, not the original predicate sweep.'),
5:('DuckDB execution and measurement','2:45–3:25',
'DuckDB runs with one thread. Each selected input is materialized and analyzed before timing. Join queries compare normal optimization with two fixed orders; Q1 has only the optimized control. Each point and control has five warmups and thirty measured executions, with randomized scheduling. The graph shows medians and bootstrap intervals. Operator profiling runs separately. All two thousand twenty-six materialized point-and-control records pass the saved result checks against the direct-input reference, with numerical tolerance. Query timing includes fetching results but excludes materialization and ANALYZE. It is not end-to-end latency.'),
6:('DuckDB runtime and observed plan changes','3:25–4:15',
'Each colored line follows one query across retained input percentages. The shaded intervals describe uncertainty in the median; the right panel enlarges the first ten percent. Beneath the curves, each query has a plan-change track. Diamonds mark join-association changes, triangles build-and-probe orientation changes, and squares physical-only changes. Stacked symbols mean dimensions changed together. These markers come from saved plans and refinement probes, not from bends in the runtime curves. Labels give selected association thresholds; exact row boundaries are in the backup notes. The search is not exhaustive. A changed plan need not produce a visible timing jump.'),
7:('Q8 changes its selected join association','4:15–4:55',
'Here are two observed Q8 associations on the same database. At one percent retention, two hundred selected parts join Lineitem early. At five percent, one thousand parts attach after the region, customer, orders and lineitem branch; Supplier and its nation form another branch. These diagrams show association, not physical build-and-probe orientation. The change motivates candidate structures to test in ORQ. It does not establish faster secure execution. The ORQ results that follow are earlier fixed-predicate measurements of canonical transfers frozen on the first of September.'),
14:('Conclusion','9:10–9:45',
'The materialized DuckDB study shows how controlled physical input sizes can change selected join structure. Separately, the earlier ORQ measurements suggest a compatible fixed transfer can reduce secure execution cost, with Q5 the strongest case. These findings motivate further tests, not a guarantee that materialization or every DuckDB plan helps ORQ. Next, validate matched inputs and exact outputs for those historical comparisons, then test new candidates with paired measurements and controlled public input sizes. That would establish when the approach helps beyond the current cases.')}


def build_notes(p,transitions,lookup):
    notes=ORIGINAL_NOTES.read_text()
    for i,(title,timing,script) in SCRIPTS.items():
        replacement=f'### Slide {i} — {title}\n\n**Target {timing}**\n\n{script}\n\n'
        notes=re.sub(rf'### Slide {i} — .*?(?=### Slide \d+ — |## Question navigation)',lambda m:replacement,notes,count=1,flags=re.S)
    notes=notes.replace('### Slide 8 — ORQ data and measurement','### Slide 8 — Earlier ORQ study: data and measurement')
    notes=notes.replace('### Slide 12 — Original and DuckDB-derived ORQ runtime','### Slide 12 — Earlier ORQ runtime: original and transferred plans')
    notes=notes.replace('the later selectivity sweep characterizes DuckDB behaviour.','the 8 September materialized-input study is a separate, later characterization.')
    source=f'Sources: [{RUN.name}](results/duckdb-tpch-selectivity/{RUN.name}/manifest.json), `points.csv`, `samples.csv`, `preparations.csv`, `transition_search.json`, and exact SQL/EXPLAIN files under `plans/` in that directory.'
    details={
3:'Q5 replaces the complete original date window with a physically materialized Orders subset. Region ASIA remains fixed. `CREATE OR REPLACE TEMP TABLE selected_orders AS SELECT * FROM controlled_orders WHERE selectivity_rank <= k; ANALYZE selected_orders;` precedes the query. The query reads selected_orders without an in-query rank predicate. The illustrated associations are the optimized common points k=1500 (1%) and k=3000 (2%). The association boundary is k=2501 → 2502 (1.667333333% → 1.668%). Trees omit build/probe orientation and output operators; retained rows are not final output cardinalities.',
4:'DuckDB 1.5.4, DBGEN SF0.1, one fixed database. Original sizes: Lineitem 600572, Orders 150000, Part 20000. Q1/Q3 vary Lineitem; Q5 Orders; Q8/Q9 Part. Seed 20260908. Rank uses row_number ordered by MD5 of seed and stable primary key, with primary-key tie breaking. k is rounded half up from target percent × original rows / 100. Nested rank prefixes are materialized; ranking is prepared once. The 31 common targets span 0.01–100%; refinements are additional measured points. Q3 keeps BUILDING and the order-date cutoff; only its ship-date condition is replaced. Q5 keeps ASIA. Q8 keeps AMERICA, its order-date range and Brazil CASE expression. Q1/Q9 have no remaining input predicates. Selected subsets do not reproduce the original predicate semantics or establish cross-engine data equivalence.',
5:'Manifest status is complete, finished 2026-09-08T21:20:20Z. One thread, five warmups, thirty unprofiled execute-through-fetchall timings per point/control; separate single EXPLAIN ANALYZE profile. Randomized point ordering and per-block control ordering. Materialization and ANALYZE, one-time rank preparation, optimizer setting changes, validation and profiling are outside query timings. Preparation CSV contains single CTAS + ANALYZE observations, not repeated end-to-end timings. Median uncertainty is a seeded 95% percentile bootstrap interval from 2000 resamples. Two fixed controls disable join_order and build_side_probe_side, constrain SQL association and orientation, and leave other optimizations active. Q1 only has optimized execution. All 2026 materialized point/control records report result_equivalent=True against the direct-input rank-filter reference. This includes common and refinement points, not 2026 distinct percentages. Numeric checks use relative and absolute tolerance 1e-8; Q3 separately checks all groups and permits valid Top-N boundary ties. Five separate canonical validations pass. This is evidence from the recorded checks, not independent proof of every query rewrite. Do not carry over the old experiment’s presenter-confirmed host or recovered memory settings to this run.',
6:'Combined plot: materialized + optimized + point_kind=common only, 31 measurements per query, sorted by actual_percent. Y is median_seconds × 1000, with median_ci_low/high × 1000. Raw samples independently reproduce each plotted median. The right panel repeats the 0–10% region at the same runtime scale. Connections are visual guides, not sampled values between observations. Tracks use all saved materialized transition-search endpoints, including refinement rows, independently of the common timing grid. Diamonds denote association changes; triangles orientation changes (which can accompany an association change); squares are reserved for physical fingerprint changes with neither association nor orientation changes. A physical fingerprint often changes automatically with either structural change and is not counted as an additional independent operator change. Exact fingerprints exclude some physical-plan details such as estimates; do not call them exhaustive optimizer decisions. Numeric labels on the chart are association thresholds rounded to three decimals. All recorded boundaries are adjacent row counts; their first observed new-plan percentage is the marker position. Unknown intervals must be shaded if future artifacts have non-adjacent boundaries. No extrapolation of thresholds to another DuckDB version or data distribution. Equal retained percentages across queries do not imply equal work. Full endpoint evidence follows below.',
7:'The two Q8 diagrams are generated from association_fingerprint at q08.materialized.k200.common.p1.optimized and q08.materialized.k1000.common.p5.optimized. Nation aliases are distinguished as region and name roles. These are binary logical association trees, not full physical operator plans or ORQ execution trees. The 1% and 5% examples have additional observed transitions between them; this pair does not claim a unique boundary or an isolated causal timing comparison. Q8 has five distinct observed associations over all materialized probes. Materialization plus ANALYZE changes both physical input and information available to planning; it does not isolate every cause of the optimizer response or prove an end-to-end speedup.',
14:'Keep future work spoken only. The completed 8 September DuckDB materialized-input study is separate from the original ORQ fixed-predicate comparison; canonical transfers were frozen 1 September. Retain the historical ORQ qualification: profiling disabled exact output validation. The new DuckDB checks do not validate those historical ORQ outputs. Safe public input-size experiments and matched ORQ output validation are needed before assigning benefits to the new candidate plans.'}
    for i,detail in details.items():
        notes=re.sub(rf'### Slide {i} detail\n.*?(?=### Slide \d+ detail|## Backup notes)',
                     lambda m:f'### Slide {i} detail\n\n{detail}\n\n{source}\n\n',notes,count=1,flags=re.S)
    notes=notes.replace('the 7 Sep sweep characterizes DuckDB behaviour and does not adaptively choose a plan inside ORQ.',
                        'the 8 Sep materialized-input study is a later characterization and does not adaptively choose a plan inside ORQ.')
    backup='''### Slide 17

Q&A

Q: Where exactly does a plan change? A: The table lists first-new-plan percentages rounded to three decimals. The endpoint ledger below provides adjacent before/after row counts, unrounded percentages, changed dimensions, plan IDs and source paths.

Q: Are all operator changes join-order changes? A: No. Association describes which relations join together. Orientation preserves child ordering and captures build/probe structure. Physical fingerprints additionally distinguish operator-tree structure and operator types. A square marks a physical-only change; an accompanying physical fingerprint change is not a separate independent event when association or orientation already changes.

Q: How many associations occur? A: Across observed optimized materialized points, Q1 has one (no joins), Q3/Q5 two each, and Q8/Q9 five each. The direct-input arm of this same controlled study has one observed association per query. This comparison is distinct from the historical 7 September predicate sweep.

Q: Did materialization speed up ORQ? A: This DuckDB experiment does not measure that. Materialization and ANALYZE are excluded from its query times, and secret ORQ validity bits do not automatically shrink physical input lengths. The ORQ slides retain the earlier canonical-transfer measurements.

Q: Are thresholds exhaustive or universal? A: No. The search probes quarter points and midpoints, then bisects observed changes within a fixed budget. It can miss changes between probes and depends on this database, settings and DuckDB version.

'''+source+'\n\n'
    notes=re.sub(r'### Slide 17\n.*?(?=### Slide 18\n)',lambda m:backup,notes,count=1,flags=re.S)
    a=notes.index('## Question navigation');b=notes.index('## Technical preparation',a)
    notes=notes[:a]+'''## Question navigation

| Topic | Slide |
|---|---:|
| Q5 materialization and join-association example | 3 |
| Input retention, replaced predicates, nested subsets | 4 |
| Repeated timings, controls, result checks, preparation exclusion | 5 |
| Combined query lines, zoom and plan-change tracks | 6 |
| Q8 association trees | 7 |
| Earlier ORQ inputs, network and measurements | 8–13 |
| Conclusion and spoken future work | 14 |
| Historical network evidence / runtime means | 15 / 16 |
| Materialized plan transitions and direct-input context | 17 |
| Secret counts versus public physical lengths | 18 |

'''+notes[b:]
    notes=notes[:notes.index('## Companion files')]
    notes+='## Exact materialized transition evidence\n\nA = association; O = orientation; P = physical-only. Physical fingerprints can also change as a consequence of A/O. All percentages use the original varied relation as denominator.\n\n'
    notes+='| Query / boundary | Before → after rows | Before → after retention (%) | Raw changed dimensions |\n|---|---:|---:|---|\n'
    for t in transitions:
        notes+=f"| {t['transition']} | {t['low_count']:,} → {t['high_count']:,} | {t['low_percent']:.9f} → {t['high_percent']:.9f} | {', '.join(t['changed_dimensions'])} |\n"
    notes+='\n### Endpoint plan identifiers and sources\n\n'
    for t in transitions:
        notes+=f"**{t['transition']}**\n\n"
        for side in ['low','high']:
            path=f"results/duckdb-tpch-selectivity/{RUN.name}/{t[side+'_plan_path']}"
            notes+=f"- [{side}: {t[side+'_point_id']}]({path}): association `{t[side+'_association_id']}`, orientation `{t[side+'_orientation_id']}`, physical `{t[side+'_physical_id']}`.\n"
        notes+='\n'
    notes+='## Companion files\n\n'
    for label,path in [('Revised PowerPoint',DECK.name),('Combined line chart','outputs/duckdb-materialized-presentation/duckdb-materialized-lines.png'),('Full-resolution vector chart','outputs/duckdb-materialized-presentation/duckdb-materialized-lines.svg'),('Machine-readable transition ledger','outputs/duckdb-materialized-presentation/transitions.json'),('Historical presentation',ORIGINAL.name),('Historical speaker notes',ORIGINAL_NOTES.name)]: notes+=f'- [{label}]({path})\n'
    spoken=notes.split('## Speaking script\n')[1].split('## Question navigation')[0]
    spoken=re.sub(r'^###.*$|^\*\*Target.*$', '',spoken,flags=re.M)
    count=len(spoken.split())
    notes=notes.replace('# ORQ evaluation — speaker notes','# ORQ evaluation — materialized-input revision: speaker notes',1)
    notes=re.sub(r'approximately [\d,]+ spoken words',f'approximately {count:,} spoken words',notes,count=1)
    NOTES.write_text(notes)
    # Embed the same per-slide script and preparation as the companion Markdown.
    for i,slide in enumerate(p.slides,1):
        if i<=14:
            speech=re.search(rf'### Slide {i} — .*?(?=### Slide \d+ — |## Question navigation)',notes,re.S).group(0)
            detail=re.search(rf'### Slide {i} detail\n.*?(?=### Slide \d+ detail|## Backup notes)',notes,re.S).group(0)
            content=speech+'\nPREPARATION — NOT ADDITIONAL NARRATION\n\n'+detail
        else:
            content=re.search(rf'### Slide {i}\n.*?(?=### Slide \d+\n|## Exact materialized)',notes,re.S).group(0)
        if i in [6,17]:content+='\n'+notes.split('## Exact materialized transition evidence')[1].split('## Companion files')[0]
        slide.notes_slide.notes_text_frame.text=content.strip()
    return count


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    before={f.name:sha(f) for f in [ORIGINAL,ORIGINAL_NOTES]}
    manifest,rows,selected,common,lookup,transitions=evidence()
    (OUT/'transitions.json').write_text(json.dumps(transitions,indent=2)+'\n')
    with (OUT/'plotted-common.csv').open('w',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(common[0]));writer.writeheader();writer.writerows(common)
    chart(common,transitions)
    p=build_slides(lookup,transitions,selected)
    word_count=build_notes(p,transitions,lookup)
    p.save(DECK)
    for f in [ORIGINAL,ORIGINAL_NOTES]:assert sha(f)==before[f.name]
    verify=Presentation(DECK);assert len(verify.slides)==18
    for i,s in enumerate(verify.slides,1):assert s.notes_slide.notes_text_frame.text.strip()
    sources={f:sha(RUN/f) for f in ['manifest.json','points.csv','samples.csv','transition_search.json']}
    (OUT/'build-manifest.json').write_text(json.dumps({'source_run':str(RUN.relative_to(ROOT)),
        'source_sha256':sources,'preserved_original_sha256':before,'deck_sha256':sha(DECK),
        'notes_sha256':sha(NOTES),'common_chart_points':len(common),'transitions':len(transitions),
        'spoken_word_count':word_count,'slides':18,'validation':'saved data checks and PowerPoint round-trip passed'},indent=2)+'\n')
    print(f'Created {DECK}\nCreated {NOTES}\n{len(common)} chart points; {len(transitions)} verified transitions; {word_count} spoken words.')


if __name__=='__main__': main()
